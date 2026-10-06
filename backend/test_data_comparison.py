from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook

from data_comparison import (
    ComparisonUploadError,
    build_pricing_scenario,
    build_results_workbook,
    compare_company_rows,
    interpret_analysis_request,
    load_analysis_run,
    parse_uploaded_table,
    parse_uploaded_dataset,
    save_analysis_run,
    summarize_results,
)


class DataComparisonTests(unittest.TestCase):
    def test_csv_upload_detects_company_and_location_columns(self):
        upload = parse_uploaded_table(
            "members.csv",
            b"Company Name,City,Postcode\nAcme Laundry Ltd,Leeds,LS1 1AA\n",
        )

        self.assertEqual(upload["company_column"], "Company Name")
        self.assertEqual(upload["rows"][0]["source_company"], "Acme Laundry Ltd")
        self.assertEqual(upload["rows"][0]["source_city"], "Leeds")
        self.assertEqual(upload["rows"][0]["source_zip"], "LS1 1AA")

    def test_upload_reports_available_columns_when_company_column_is_unknown(self):
        with self.assertRaisesRegex(ComparisonUploadError, "Available columns: Member ID, Trading title"):
            parse_uploaded_table(
                "members.csv",
                b"Member ID,Trading title\n1,Acme Laundry\n",
            )

    def test_matching_separates_exact_possible_and_unmatched_rows(self):
        companies = [
            {
                "primary_key": "C1",
                "company": "Acme Laundry Services LLC",
                "parent_org": "",
                "city": "Leeds",
                "state": "",
                "zip_code": "LS1 1AA",
                "price_list": "D",
            },
            {
                "primary_key": "C2",
                "company": "Bright Textile Care",
                "parent_org": "",
                "city": "York",
                "state": "",
                "zip_code": "",
                "price_list": "A",
            },
        ]
        rows = [
            {"source_row": 2, "source_company": "Acme Laundry Services Ltd", "source_city": "Leeds", "source_state": "", "source_zip": "LS1 1AA"},
            {"source_row": 3, "source_company": "Bright Tex Care", "source_city": "", "source_state": "", "source_zip": ""},
            {"source_row": 4, "source_company": "Completely Different", "source_city": "", "source_state": "", "source_zip": ""},
        ]

        results = compare_company_rows(rows, companies)

        self.assertEqual(results[0]["match_status"], "Matched")
        self.assertTrue(results[0]["is_price_list_d"])
        self.assertEqual(results[1]["match_status"], "Possible match")
        self.assertEqual(results[2]["match_status"], "No match")

    def test_ambiguous_exact_name_requires_review(self):
        rows = [{"source_company": "Acme Laundry", "source_city": "", "source_state": "", "source_zip": ""}]
        companies = [
            {"primary_key": "C1", "company": "Acme Laundry Ltd", "parent_org": "", "price_list": "D"},
            {"primary_key": "C2", "company": "Acme Laundry LLC", "parent_org": "", "price_list": "D"},
        ]

        result = compare_company_rows(rows, companies)[0]

        self.assertEqual(result["match_status"], "Possible match")
        self.assertFalse(result["is_price_list_d"])
        self.assertIn("multiple close", result["reason"])

    def test_order_price_list_is_used_as_fallback(self):
        rows = [{"source_company": "Acme Laundry", "source_city": "", "source_state": "", "source_zip": ""}]
        companies = [{"primary_key": "C1", "company": "Acme Laundry", "parent_org": "", "price_list": ""}]

        result = compare_company_rows(
            rows,
            companies,
            price_list_by_key={"C1": "D"},
            customer_metrics_by_key={"C1": {"order_count": 4, "total_spend": 1234.5}},
        )[0]

        self.assertTrue(result["is_price_list_d"])
        self.assertEqual(result["order_count"], 4)
        self.assertEqual(result["total_spend"], 1234.5)

    def test_plain_language_request_recognizes_price_list_and_spend(self):
        request = interpret_analysis_request(
            "Identify the customer and associated price list, and include how much they have spent with us all time."
        )

        self.assertTrue(request["match_companies"])
        self.assertTrue(request["include_price_list"])
        self.assertTrue(request["include_spend"])

    def test_pricing_request_is_detected_and_grouped(self):
        dataset = parse_uploaded_dataset(
            "mats.csv",
            (
                "Price List,Mat Size,Sales Amount\n"
                "A,3x5,100\n"
                "A,3x5,50\n"
                "B,4x6,200\n"
            ).encode(),
        )

        result = build_pricing_scenario(
            dataset,
            "What would sales look like if we increase prices by 5% and break it down by price list and mat size?",
        )

        self.assertEqual(result["analysis_type"], "pricing")
        self.assertEqual(result["summary"]["current_sales"], 350.0)
        self.assertEqual(result["summary"]["projected_sales"], 367.5)
        self.assertEqual(result["summary"]["change"], 17.5)
        self.assertEqual(result["results"][0]["current_sales"], 150.0)
        self.assertEqual(result["results"][0]["projected_sales"], 157.5)

    def test_pricing_scenario_supports_unit_value_times_quantity(self):
        dataset = parse_uploaded_dataset(
            "mats.csv",
            b"Price List,Mat Size,Unit Price,Quantity\nD,3x5,10,4\n",
        )

        result = build_pricing_scenario(dataset, "Decrease prices by 10% by price list and mat size")

        self.assertEqual(result["value_mode"], "unit")
        self.assertEqual(result["summary"]["current_sales"], 40.0)
        self.assertEqual(result["summary"]["projected_sales"], 36.0)
        self.assertEqual(result["summary"]["change"], -4.0)

    def test_pricing_scenario_reports_non_numeric_rows(self):
        dataset = parse_uploaded_dataset(
            "mats.csv",
            b"Price List,Mat Size,Sales Amount\nA,3x5,100\nB,4x6,n/a\n",
        )

        result = build_pricing_scenario(dataset, "Increase prices by 5%")

        self.assertEqual(result["summary"]["included_rows"], 1)
        self.assertEqual(result["summary"]["excluded_rows"], 1)
        self.assertEqual(result["excluded_rows"], [3])

    def test_pricing_workbook_has_summary_and_detail(self):
        dataset = parse_uploaded_dataset(
            "mats.csv",
            b"Price List,Mat Size,Sales Amount\nA,3x5,100\n",
        )
        payload = build_pricing_scenario(dataset, "Increase prices by 5% by price list and mat size")
        payload["filename"] = "mats.csv"

        workbook = load_workbook(BytesIO(build_results_workbook(payload)), data_only=True)

        self.assertEqual(workbook.sheetnames, ["Pricing scenario", "Detail"])
        self.assertEqual(workbook["Pricing scenario"]["B9"].value, 105)
        self.assertEqual(workbook["Detail"]["G2"].value, 105)

    def test_pricing_scenario_requires_percentage(self):
        dataset = parse_uploaded_dataset("mats.csv", b"Sales Amount\n100\n")

        with self.assertRaisesRegex(ComparisonUploadError, "percentage"):
            build_pricing_scenario(dataset, "Show a pricing scenario")

    def test_saved_run_and_excel_export_include_all_results(self):
        results = [{
            "source_row": 2, "source_company": "Acme Laundry", "source_city": "Leeds", "source_state": "",
            "source_zip": "LS1", "match_status": "Matched", "matched_company": "Acme Laundry",
            "customer_primary_key": "C1", "matched_city": "Leeds", "matched_state": "", "matched_zip": "LS1",
            "price_list": "D", "is_price_list_d": True, "confidence": 100, "reason": "exact normalized name",
        }]
        payload = {"filename": "members.csv", "company_column": "Company", "results": results, "summary": summarize_results(results), "owner": "tester"}
        with TemporaryDirectory() as temp_dir:
            token = save_analysis_run(payload, run_dir=Path(temp_dir))
            loaded = load_analysis_run(token, run_dir=Path(temp_dir))

        workbook = load_workbook(BytesIO(build_results_workbook(loaded)), data_only=True)
        self.assertEqual(workbook["Comparison"]["B10"].value, "Acme Laundry")
        self.assertEqual(workbook["Comparison"]["M10"].value, "Yes")

    def test_excel_export_treats_uploaded_formulas_as_text(self):
        results = [{
            "source_row": 2, "source_company": "=HYPERLINK(\"bad\")", "match_status": "No match",
            "is_price_list_d": False, "confidence": 0, "reason": "No credible company match",
        }]
        payload = {"filename": "members.csv", "company_column": "Company", "results": results, "summary": summarize_results(results)}

        workbook = load_workbook(BytesIO(build_results_workbook(payload)), data_only=False)

        self.assertEqual(workbook["Comparison"]["B10"].data_type, "s")


if __name__ == "__main__":
    unittest.main()
