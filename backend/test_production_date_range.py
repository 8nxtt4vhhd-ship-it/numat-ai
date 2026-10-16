import unittest
from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook

import main
from main import render_production_date_control, resolve_production_date_range
from production import build_production_kpi_payload


class ProductionDateRangeTests(unittest.TestCase):
    def test_cycle_detail_replaces_saved_daily_press_total(self):
        result = {
            "status": "ok",
            "production_rows": [{
                "date": "2026-10-15", "press_throughput": 2429,
                "ff1_lf": 300, "ff2_lf": 1230, "ff3_lf": 899,
            }],
            "operator_rows": [], "plant_operator_rows": [],
        }
        cycles = [
            {"date": "2026-10-15", "machine": "FF1", "linear_feet": 299, "ff1_lf": 299, "ff2_lf": None, "ff3_lf": None},
            {"date": "2026-10-15", "machine": "FF2", "linear_feet": 1241, "ff1_lf": None, "ff2_lf": 1241, "ff3_lf": None},
            {"date": "2026-10-15", "machine": "FF3", "linear_feet": 912, "ff1_lf": None, "ff2_lf": None, "ff3_lf": 912},
        ]

        payload = build_production_kpi_payload(result, period_start="2026-10-15", period_end="2026-10-15", cycles=cycles)
        latest = payload["latest"]

        self.assertEqual(latest["press_throughput"], 2452)
        self.assertEqual(latest["press_throughput_daily_summary"], 2429)
        self.assertEqual((latest["ff1_lf"], latest["ff2_lf"], latest["ff3_lf"]), (299, 1241, 912))
        self.assertEqual(payload["summary"]["average_press_throughput"], 2452)

    def test_accepts_us_formatted_dates(self):
        start, end = resolve_production_date_range("10/01/2026", "10/06/2026")

        self.assertEqual(start, "2026-10-01")
        self.assertEqual(end, "2026-10-06")

    def test_control_displays_us_formatted_dates(self):
        html = render_production_date_control("2026-10-01", "2026-10-06")

        self.assertIn('value="10/01/2026"', html)
        self.assertIn('value="10/06/2026"', html)
        self.assertIn('placeholder="MM/DD/YYYY"', html)
        self.assertIn('name="refresh" value="1"', html)

    def test_default_dashboard_open_uses_last_cached_date_range(self):
        payload = {
            "status": "ok",
            "synced_at": "2026-10-07 04:00:00",
            "rows": [],
            "operator_summary": [],
            "latest": {},
            "summary": {},
        }
        detail = {"status": "ok", "clocking_count": 1, "summary": {}, "department_rows": [], "operator_rows": []}
        with (
            patch.object(main, "get_last_productivity_detail_range", return_value=("2026-10-01", "2026-10-07")),
            patch.object(main, "fetch_production_analysis_data", return_value={"status": "ok"}),
            patch.object(main, "build_production_kpi_payload", return_value=payload),
            patch.object(main, "fetch_productivity_detail", return_value=detail) as detail_fetch,
            patch.object(main, "build_current_month_production_payload", return_value={"summary": {}}),
        ):
            response = main.production_analysis_view()

        html = response.body.decode() if hasattr(response, "body") else str(response)
        self.assertIn('value="10/01/2026"', html)
        self.assertIn('value="10/07/2026"', html)
        self.assertEqual(detail_fetch.call_args.kwargs["period_start"], "2026-10-01")
        self.assertEqual(detail_fetch.call_args.kwargs["period_end"], "2026-10-07")
        self.assertIn("Export press analysis", html)
        self.assertIn("/production-analysis/export-press-analysis?start=2026-10-01&end=2026-10-07", html)

    def test_productivity_detail_page_shows_exact_calculation(self):
        detail = {
            "status": "ok",
            "booking_count": 20,
            "clocking_count": 10,
            "summary": {
                "booked_hours": 91.2,
                "clocked_hours": 100,
                "unbooked_hours": 8.8,
                "unmatched_booked_hours": 0,
                "manager_support_hours": 2,
                "manager_clocked_hours": 31.3,
                "labour_clocked_hours": 131.3,
                "plant_productivity": 91.2,
            },
            "department_rows": [],
            "operator_rows": [],
        }
        with patch.object(main, "fetch_productivity_detail", return_value=detail):
            response = main.production_productivity_detail_view("10/01/2026", "10/06/2026")

        html = response.body.decode() if hasattr(response, "body") else str(response)
        self.assertIn("Productive booked hours", html)
        self.assertIn("Production operator clocked hours", html)
        self.assertIn("manager clockings excluded from productivity", html)
        self.assertIn("Total paid plant hours", html)
        self.assertIn("31.3 manager hrs", html)
        self.assertIn("91.2%", html)
        self.assertIn("10/01/2026 to 10/06/2026", html)

    def test_operator_table_links_to_selected_period_detail(self):
        html = main.render_operator_analysis(
            [{"name": "Zach Holmes", "days": 1, "productivity": 120, "clocked_hours": 8, "booked_hours": 9.6, "unbooked_hours": 0}],
            period_start="2026-10-01",
            period_end="2026-10-06",
        )

        self.assertIn("/production-analysis/operator-detail?", html)
        self.assertIn("Zach+Holmes", html)
        self.assertIn("View bookings", html)

    def test_overbooked_report_lists_only_operator_days_above_clocked_time(self):
        detail = {
            "status": "ok",
            "time_bookings": [
                {"operator": "Zach Holmes", "operator_key": "zachholmes", "date": "2026-10-05", "hours": 5.0},
                {"operator": "Zach Holmes", "operator_key": "zachholmes", "date": "2026-10-05", "hours": 4.0},
                {"operator": "Amy Smith", "operator_key": "amysmith", "date": "2026-10-05", "hours": 7.0},
                {"operator": "Lois", "operator_key": "lois", "date": "2026-10-05", "hours": 10.0, "manager_support": True},
            ],
            "clockings": [
                {"operator": "Zach Holmes", "operator_key": "zachholmes", "date": "2026-10-05", "clocked_hours": 8.0},
                {"operator": "Amy Smith", "operator_key": "amysmith", "date": "2026-10-05", "clocked_hours": 8.0},
                {"operator": "Lois", "operator_key": "lois", "date": "2026-10-05", "clocked_hours": 8.0, "manager": True},
            ],
        }
        with patch.object(main, "fetch_productivity_detail", return_value=detail):
            response = main.production_overbooked_report("10/05/2026", "10/05/2026")

        html = response.body.decode() if hasattr(response, "body") else str(response)
        self.assertIn("Zach Holmes", html)
        self.assertIn("+1.00 hrs", html)
        self.assertIn("112.5%", html)
        self.assertIn("Review bookings", html)
        self.assertNotIn("Amy Smith", html)
        self.assertNotIn(">Lois<", html)

    def test_operator_detail_summarizes_booked_time_by_department(self):
        detail = {
            "status": "ok",
            "time_bookings": [
                {"record_id": "1", "operator": "Zach Holmes", "operator_key": "zachholmes", "date": "2026-10-05", "department": "Press", "hours": 3.0},
                {"record_id": "2", "operator": "Zach Holmes", "operator_key": "zachholmes", "date": "2026-10-05", "department": "Press", "hours": 2.0},
                {"record_id": "3", "operator": "Zach Holmes", "operator_key": "zachholmes", "date": "2026-10-05", "department": "Engineering", "hours": 5.0},
            ],
            "clockings": [],
            "operator_rows": [{"name": "Zach Holmes", "booked_hours": 10.0, "clocked_hours": 8.0, "productivity": 125.0}],
        }
        with (
            patch.object(main, "fetch_productivity_detail", return_value=detail),
            patch.object(main, "get_current_session_user", return_value={"role": "user"}),
        ):
            response = main.production_operator_detail_view("Zach Holmes", "10/05/2026", "10/05/2026")

        html = response.body.decode() if hasattr(response, "body") else str(response)
        self.assertIn("Time by department", html)
        self.assertIn("Engineering", html)
        self.assertIn("Press", html)
        self.assertIn("5.00 hrs", html)
        self.assertIn("50.0%", html)
        self.assertIn("10.00 hrs", html)
        self.assertIn("Export Excel report", html)
        self.assertIn("/production-analysis/operator-detail/export?", html)

    def test_operator_excel_report_highlights_daily_lost_time(self):
        detail = {
            "status": "ok",
            "time_bookings": [
                {"record_id": "1", "operator": "James Frieson", "operator_key": "jamesfrieson", "date": "2026-10-05", "department": "Press", "hours": 5.4, "start_time": "10/05/2026 08:00:00", "finish_time": "10/05/2026 13:24:00"},
                {"record_id": "2", "operator": "James Frieson", "operator_key": "jamesfrieson", "date": "2026-10-06", "department": "Engineering", "hours": 8.0, "start_time": "10/06/2026 08:00:00", "finish_time": "10/06/2026 16:00:00"},
            ],
            "clockings": [
                {"record_id": "10", "operator": "James Frieson", "operator_key": "jamesfrieson", "date": "2026-10-05", "clocked_hours": 8.0, "clock_in": "10/05/2026 08:00:00", "clock_out": "10/05/2026 16:00:00"},
                {"record_id": "11", "operator": "James Frieson", "operator_key": "jamesfrieson", "date": "2026-10-06", "clocked_hours": 8.0, "clock_in": "10/06/2026 08:00:00", "clock_out": "10/06/2026 16:00:00"},
            ],
            "operator_rows": [{"name": "James Frieson", "booked_hours": 13.4, "clocked_hours": 16.0, "productivity": 83.75}],
        }
        with patch.object(main, "fetch_productivity_detail", return_value=detail):
            response = main.export_production_operator_detail("James Frieson", "10/05/2026", "10/06/2026")

        workbook = load_workbook(BytesIO(response.body), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Summary", "Departments", "Time Bookings", "Clockings"])
        summary = workbook["Summary"]
        self.assertEqual(summary["C5"].value, "Lost / unbooked time")
        self.assertAlmostEqual(summary["C6"].value, 2.6)
        daily_rows = list(summary.iter_rows(min_row=12, values_only=True))
        self.assertEqual(daily_rows[0][3], 2.6)
        self.assertEqual(daily_rows[0][6], "Lost time")
        self.assertEqual(daily_rows[1][3], 0)
        self.assertEqual(workbook["Departments"]["A6"].value, "Engineering")
        self.assertIn("James-Frieson-productivity-2026-10-05-to-2026-10-06.xlsx", response.headers["content-disposition"])

    def test_press_analysis_export_includes_dashboard_and_cycle_trace(self):
        production_result = {"status": "ok"}
        payload = {
            "rows": [
                {
                    "date": "2026-10-08",
                    "ff1_lf": 1000,
                    "ff2_lf": 1200,
                    "ff3_lf": 800,
                    "ff1_mats": 10,
                    "ff2_mats": 12,
                    "ff3_mats": 8,
                    "ff1_cycles": 3,
                    "ff2_cycles": 4,
                    "ff3_cycles": 2,
                    "ff1_revenue": 1000,
                    "ff2_revenue": 1200,
                    "ff3_revenue": 800,
                }
            ]
        }
        detail = {
            "cycles": [
                {"date": "2026-10-08", "start_time": "10/08/2026 08:00:00", "finish_time": "10/08/2026 08:30:00", "order_reference": "3736", "cycle_number": "3736-1", "machine": "FF1", "operator": "Kenley Pike", "mat_count": 2, "linear_feet": 100, "ff1_lf": 100, "repair_value": 250, "record_id": "1"},
                {"date": "2026-10-08", "start_time": "10/08/2026 09:00:00", "finish_time": "10/08/2026 09:30:00", "order_reference": "3736", "cycle_number": "3736-2", "machine": "FF2", "operator": "Cedrick Brown", "mat_count": 2, "linear_feet": 120, "ff2_lf": 120, "repair_value": 300, "record_id": "2"},
            ]
        }
        with (
            patch.object(main, "fetch_production_analysis_data", return_value=production_result),
            patch.object(main, "build_production_kpi_payload", return_value=payload),
            patch.object(main, "fetch_productivity_detail", return_value=detail),
        ):
            response = main.export_production_press_analysis(start="10/08/2026", end="10/08/2026")

        workbook = load_workbook(BytesIO(response.body), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Summary", "Daily Press", "Cycle Detail", "Shared Orders"])
        self.assertEqual(workbook["Daily Press"]["E6"].value, "=SUM(B6:D6)")
        self.assertEqual(workbook["Cycle Detail"]["D6"].value, "3736")
        self.assertEqual(workbook["Cycle Detail"]["G7"].value, "Cedrick Brown")
        self.assertEqual(workbook["Shared Orders"]["B6"].value, "3736")
        self.assertIn("Cedrick Brown", workbook["Shared Orders"]["C6"].value)
        self.assertIn("Kenley Pike", workbook["Shared Orders"]["C6"].value)
        self.assertIn("press-analysis-2026-10-08-to-2026-10-08.xlsx", response.headers["content-disposition"])

    def test_time_booking_update_writes_verified_record_and_clears_cache(self):
        detail = {
            "time_bookings": [
                {
                    "record_id": "123",
                    "operator_key": "zachholmes",
                    "department": "Grind",
                    "start_time": "10/05/2026 08:00:00",
                    "finish_time": "10/05/2026 10:00:00",
                }
            ]
        }
        with (
            patch.object(main, "get_current_session_user", return_value={"username": "admin", "role": "admin"}),
            patch.object(main, "fetch_productivity_detail", return_value=detail),
            patch.object(main, "update_layout_record", return_value={"status": "ok"}) as update,
            patch.object(main, "clear_productivity_detail_cache") as clear_cache,
            patch.object(main, "clear_weekly_kpi_payload_cache") as clear_weekly_cache,
            patch.object(main, "record_audit_event"),
        ):
            response = main.update_production_time_booking(
                record_id="123",
                operator="Zach Holmes",
                start="10/01/2026",
                end="10/06/2026",
                department="Engineering",
                booking_start="2026-10-05T08:30:00",
                booking_finish="2026-10-05T09:30:00",
                notes="Corrected",
            )

        self.assertEqual(response.status_code, 303)
        update.assert_called_once_with(
            "ai_Time Bookings",
            "123",
            {
                "Department": "Engineering",
                "Start Time": "10/05/2026 08:30:00",
                "Finish Time": "10/05/2026 09:30:00",
                "Notes": "Corrected",
            },
        )
        clear_cache.assert_called_once_with()
        clear_weekly_cache.assert_called_once_with()


if __name__ == "__main__":
    unittest.main()
