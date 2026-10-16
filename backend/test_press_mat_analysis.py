from io import BytesIO
import unittest

from openpyxl import load_workbook

from production_reports import build_press_analysis_workbook


class PressMatAnalysisTests(unittest.TestCase):
    def test_press_export_falls_back_to_cycle_mat_summary(self):
        cycles = [
            {"date": "2026-10-15", "cycle_number": "3681-1", "machine": "FF2", "operator": "Kenley Pike", "mat_count": 5, "linear_feet": 80, "repair_value": 200},
            {"date": "2026-10-15", "cycle_number": "3681-2", "machine": "FF2", "operator": "Kenley Pike", "mat_count": 3, "linear_feet": 48, "repair_value": 120},
        ]

        content = build_press_analysis_workbook("2026-10-15", "2026-10-15", [], cycles)
        workbook = load_workbook(BytesIO(content), data_only=False)
        sheet = workbook["Mat Analysis"]

        self.assertEqual(sheet["A6"].value, "Kenley Pike")
        self.assertEqual(sheet["B6"].value, 8)
        self.assertEqual(sheet["C6"].value, 128)
        self.assertEqual(sheet["D6"].value, 16)
        self.assertEqual(sheet["G6"].value, 320)
        self.assertEqual(sheet["J6"].value, "Cycle summary only")

    def test_press_export_includes_mat_analysis_by_operator(self):
        cycles = [
            {"date": "2026-10-15", "cycle_number": "3681-1", "machine": "FF2", "operator": "Kenley Pike", "linear_feet": 40, "repair_value": 100},
            {"date": "2026-10-15", "cycle_number": "3681-2", "machine": "FF3", "operator": "Xavier Slaughter", "linear_feet": 20, "repair_value": 30},
        ]
        mats = [
            {"press_date": "2026-10-15", "order_reference": "3681", "cycle_number": "3681-1", "mat_number": "3681-1", "mat_size": "4x12", "mat_type": "Standard", "mat_outcome": "Repair", "square_feet": 48, "linear_feet": 32, "repair_value": 90, "repair_count": 4, "record_id": "1"},
            {"press_date": "2026-10-15", "order_reference": "3681", "cycle_number": "3681-2", "mat_number": "3681-2", "mat_size": "3x5", "mat_type": "Standard", "mat_outcome": "Repair", "square_feet": 15, "linear_feet": 16, "repair_value": 30, "repair_count": 2, "record_id": "2"},
        ]

        content = build_press_analysis_workbook("2026-10-15", "2026-10-15", [], cycles, mats)
        workbook = load_workbook(BytesIO(content), data_only=False)

        self.assertIn("Mat Analysis", workbook.sheetnames)
        sheet = workbook["Mat Analysis"]
        self.assertEqual(sheet["A6"].value, "Kenley Pike")
        self.assertEqual(sheet["D6"].value, 32)
        self.assertEqual(sheet["G6"].value, 90)
        self.assertIn("4x12", sheet["J6"].value)
        self.assertEqual(sheet.cell(11, 5).value, "Kenley Pike")


if __name__ == "__main__":
    unittest.main()
