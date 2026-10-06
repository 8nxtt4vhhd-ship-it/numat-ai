import csv
from collections import defaultdict
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from difflib import SequenceMatcher, get_close_matches
from io import BytesIO, StringIO
import json
from pathlib import Path
import re
import secrets
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_UPLOAD_ROWS = 5000
RUN_RETENTION_HOURS = 24
DEFAULT_RUN_DIR = Path(__file__).resolve().parent / "data" / "comparison_runs"

COMPANY_HEADER_ALIASES = (
    "company",
    "company name",
    "business",
    "business name",
    "organisation",
    "organisation name",
    "organization",
    "organization name",
    "member",
    "member name",
    "laundry",
    "laundry name",
    "customer",
    "customer name",
)

LOCATION_HEADER_ALIASES = {
    "city": ("city", "town"),
    "state": ("state", "province", "county"),
    "zip_code": ("postcode", "postal code", "zip code", "zipcode", "zip"),
}

COMPANY_SUFFIXES = {
    "co", "company", "corp", "corporation", "inc", "incorporated", "limited", "ltd",
    "llc", "llp", "lp", "plc", "group", "holdings", "the",
}

VALUE_COLUMN_ALIASES = (
    "sales value", "sales amount", "net sales", "revenue", "line total", "extended price",
    "extended value", "invoice value", "repair value", "amount", "total value", "sales",
    "unit price", "selling price", "price", "cost",
)
QUANTITY_COLUMN_ALIASES = ("quantity", "qty", "units", "number of mats", "mat quantity", "count")
PRICE_LIST_COLUMN_ALIASES = ("price list", "pricelist", "price band", "pricing group")
MAT_SIZE_COLUMN_ALIASES = ("mat size", "size", "mat dimensions", "dimensions", "product size")


class ComparisonUploadError(ValueError):
    pass


def normalize_header(value):
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def normalize_company_name(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char)).casefold()
    tokens = re.findall(r"[a-z0-9]+", text)
    while tokens and tokens[-1] in COMPANY_SUFFIXES:
        tokens.pop()
    return " ".join(token for token in tokens if token != "the")


def normalize_location(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char)).casefold()
    return re.sub(r"[^a-z0-9]", "", text)


def _unique_headers(values):
    headers = []
    seen = {}
    for index, value in enumerate(values, start=1):
        base = str(value or "").strip() or f"Column {index}"
        count = seen.get(base.casefold(), 0) + 1
        seen[base.casefold()] = count
        headers.append(base if count == 1 else f"{base} ({count})")
    return headers


def _rows_from_csv(content):
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [list(row) for row in csv.reader(StringIO(text), dialect)]


def _rows_from_xlsx(content):
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise ComparisonUploadError("The Excel workbook could not be read.") from error
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        return [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def parse_uploaded_dataset(filename, content):
    filename = str(filename or "").strip()
    suffix = Path(filename).suffix.casefold()
    if suffix not in {".csv", ".xlsx"}:
        raise ComparisonUploadError("Please upload a .xlsx or .csv file.")
    if not content:
        raise ComparisonUploadError("The uploaded file is empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ComparisonUploadError("The uploaded file is larger than 10 MB.")

    raw_rows = _rows_from_csv(content) if suffix == ".csv" else _rows_from_xlsx(content)
    header_index = next(
        (index for index, row in enumerate(raw_rows[:25]) if any(str(value or "").strip() for value in row)),
        None,
    )
    if header_index is None:
        raise ComparisonUploadError("No rows were found in the uploaded file.")

    headers = _unique_headers(raw_rows[header_index])
    rows = []
    for raw_index, raw_row in enumerate(raw_rows[header_index + 1:], start=header_index + 2):
        values = list(raw_row) + [None] * max(0, len(headers) - len(raw_row))
        row = {header: values[index] for index, header in enumerate(headers)}
        if not any(value not in {None, ""} for value in row.values()):
            continue
        rows.append({"source_row": raw_index, "values": row})
        if len(rows) > MAX_UPLOAD_ROWS:
            raise ComparisonUploadError(f"The upload contains more than {MAX_UPLOAD_ROWS:,} data rows.")

    if not rows:
        raise ComparisonUploadError("No data rows were found beneath the header row.")
    return {"filename": filename, "headers": headers, "rows": rows}


def parse_uploaded_table(filename, content, company_column=""):
    dataset = parse_uploaded_dataset(filename, content)
    headers = dataset["headers"]
    normalized_headers = {normalize_header(header): header for header in headers}
    requested_column = normalize_header(company_column)
    if requested_column:
        selected_company_column = normalized_headers.get(requested_column)
        if not selected_company_column:
            raise ComparisonUploadError(
                f"Company column '{company_column}' was not found. Available columns: {', '.join(headers)}."
            )
    else:
        selected_company_column = next(
            (normalized_headers[alias] for alias in COMPANY_HEADER_ALIASES if alias in normalized_headers),
            None,
        )
        if not selected_company_column:
            raise ComparisonUploadError(
                f"Choose the company-name column. Available columns: {', '.join(headers)}."
            )

    location_columns = {}
    for location_key, aliases in LOCATION_HEADER_ALIASES.items():
        location_columns[location_key] = next(
            (normalized_headers[alias] for alias in aliases if alias in normalized_headers),
            "",
        )

    rows = []
    for source in dataset["rows"]:
        raw_index = source["source_row"]
        row = source["values"]
        company_name = str(row.get(selected_company_column) or "").strip()
        if not company_name:
            continue
        rows.append({
            "source_row": raw_index,
            "source_company": company_name,
            "source_city": str(row.get(location_columns["city"]) or "").strip() if location_columns["city"] else "",
            "source_state": str(row.get(location_columns["state"]) or "").strip() if location_columns["state"] else "",
            "source_zip": str(row.get(location_columns["zip_code"]) or "").strip() if location_columns["zip_code"] else "",
        })
    if not rows:
        raise ComparisonUploadError("No company names were found in the selected column.")
    return {
        "filename": dataset["filename"],
        "headers": headers,
        "company_column": selected_company_column,
        "rows": rows,
    }


def _company_aliases(company):
    aliases = [company.get("company"), company.get("parent_org")]
    return sorted({normalize_company_name(alias) for alias in aliases if normalize_company_name(alias)})


def _location_bonus(source, company):
    comparisons = [
        (source.get("source_city"), company.get("city"), 4),
        (source.get("source_state"), company.get("state"), 2),
        (source.get("source_zip"), company.get("zip_code"), 5),
    ]
    bonus = 0
    reasons = []
    for source_value, company_value, points in comparisons:
        if source_value and company_value and normalize_location(source_value) == normalize_location(company_value):
            bonus += points
            reasons.append("matching postcode" if points == 5 else f"matching {('city' if points == 4 else 'state')}")
    return bonus, reasons


def _name_score(source_name, alias):
    if source_name == alias:
        return 100
    source_tokens = set(source_name.split())
    alias_tokens = set(alias.split())
    token_score = SequenceMatcher(None, " ".join(sorted(source_tokens)), " ".join(sorted(alias_tokens))).ratio()
    sequence_score = SequenceMatcher(None, source_name, alias).ratio()
    containment_score = 0.0
    if min(len(source_name), len(alias)) >= 6 and (source_name in alias or alias in source_name):
        containment_score = 0.91
    return round(max(token_score, sequence_score, containment_score) * 100)


def interpret_analysis_request(request_text):
    request_text = str(request_text or "").strip()
    normalized = normalize_header(request_text)
    percentage_change = parse_percentage_change(normalized)
    pricing_terms = any(term in normalized for term in ("increase", "decrease", "pricing", "price change", "what sales", "scenario"))
    return {
        "text": request_text,
        "analysis_type": "pricing" if pricing_terms else "company_match",
        "percentage_change": percentage_change,
        "match_companies": True,
        "include_price_list": not normalized or "price" in normalized or "customer" in normalized,
        "include_spend": any(
            phrase in normalized
            for phrase in ("spend", "spent", "sales", "revenue", "order value", "all time", "lifetime")
        ),
    }


def parse_percentage_change(request_text):
    normalized = str(request_text or "").casefold()
    patterns = (
        r"(?:increase|raise|up)\D{0,20}(\d+(?:\.\d+)?)\s*%",
        r"(\d+(?:\.\d+)?)\s*%\D{0,20}(?:increase|raise|higher|uplift)",
        r"(?:decrease|reduce|lower|down)\D{0,20}(\d+(?:\.\d+)?)\s*%",
        r"(\d+(?:\.\d+)?)\s*%\D{0,20}(?:decrease|reduction|lower|cut)",
    )
    for index, pattern in enumerate(patterns):
        match = re.search(pattern, normalized)
        if match:
            value = float(match.group(1))
            return -value if index >= 2 else value
    return None


def _resolve_column(headers, requested, aliases, label, required=False):
    normalized_headers = {normalize_header(header): header for header in headers}
    requested_normalized = normalize_header(requested)
    if requested_normalized:
        resolved = normalized_headers.get(requested_normalized)
        if not resolved:
            raise ComparisonUploadError(
                f"{label} column '{requested}' was not found. Available columns: {', '.join(headers)}."
            )
        return resolved
    resolved = next((normalized_headers[alias] for alias in aliases if alias in normalized_headers), "")
    if required and not resolved:
        raise ComparisonUploadError(
            f"Choose the {label.lower()} column. Available columns: {', '.join(headers)}."
        )
    return resolved


def _decimal_value(value):
    if value in {None, ""}:
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value))
        except InvalidOperation:
            return None
    text = str(value).strip()
    negative = text.startswith("(") and text.endswith(")")
    text = re.sub(r"[^0-9.\-]", "", text)
    if not text or text in {"-", ".", "-."}:
        return None
    try:
        result = Decimal(text)
    except InvalidOperation:
        return None
    return -result if negative and result > 0 else result


def build_pricing_scenario(
    dataset,
    request_text,
    *,
    value_column="",
    quantity_column="",
    price_list_column="",
    mat_size_column="",
    value_mode="auto",
):
    headers = dataset.get("headers") or []
    change_percent = parse_percentage_change(request_text)
    if change_percent is None:
        raise ComparisonUploadError("Include a percentage increase or decrease in the request, such as 'increase prices by 5%'.")
    if change_percent <= -100:
        raise ComparisonUploadError("The price decrease must be less than 100%.")

    selected_value = _resolve_column(headers, value_column, VALUE_COLUMN_ALIASES, "Value", required=True)
    selected_quantity = _resolve_column(headers, quantity_column, QUANTITY_COLUMN_ALIASES, "Quantity")
    selected_price_list = _resolve_column(headers, price_list_column, PRICE_LIST_COLUMN_ALIASES, "Price list")
    selected_mat_size = _resolve_column(headers, mat_size_column, MAT_SIZE_COLUMN_ALIASES, "Mat size")

    mode = str(value_mode or "auto").strip().casefold()
    if mode not in {"auto", "extended", "unit"}:
        raise ComparisonUploadError("Value mode must be auto, extended or unit.")
    if mode == "auto":
        value_header = normalize_header(selected_value)
        mode = "unit" if selected_quantity and any(term in value_header for term in ("unit", "each", "price", "cost")) else "extended"
    if mode == "unit" and not selected_quantity:
        raise ComparisonUploadError("A quantity column is required when the value column contains a unit price or unit cost.")

    multiplier = Decimal("1") + (Decimal(str(change_percent)) / Decimal("100"))
    groups = defaultdict(lambda: {"row_count": 0, "quantity": Decimal("0"), "current_sales": Decimal("0")})
    details = []
    excluded_rows = []
    for source in dataset.get("rows") or []:
        values = source["values"]
        value = _decimal_value(values.get(selected_value))
        quantity = _decimal_value(values.get(selected_quantity)) if selected_quantity else Decimal("1")
        if value is None or quantity is None:
            excluded_rows.append(source["source_row"])
            continue
        current_sales = value * quantity if mode == "unit" else value
        projected_sales = (current_sales * multiplier).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        current_sales = current_sales.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        price_list = str(values.get(selected_price_list) or "Unspecified").strip() if selected_price_list else "All"
        mat_size = str(values.get(selected_mat_size) or "Unspecified").strip() if selected_mat_size else "All"
        group = groups[(price_list, mat_size)]
        group["row_count"] += 1
        group["quantity"] += quantity
        group["current_sales"] += current_sales
        details.append({
            "source_row": source["source_row"],
            "price_list": price_list,
            "mat_size": mat_size,
            "quantity": float(quantity),
            "input_value": float(value),
            "current_sales": float(current_sales),
            "projected_sales": float(projected_sales),
            "change": float(projected_sales - current_sales),
        })

    if not details:
        raise ComparisonUploadError(f"No numeric values were found in the '{selected_value}' column.")

    results = []
    for (price_list, mat_size), group in groups.items():
        current_sales = group["current_sales"].quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        projected_sales = (current_sales * multiplier).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        results.append({
            "price_list": price_list,
            "mat_size": mat_size,
            "row_count": group["row_count"],
            "quantity": float(group["quantity"]),
            "current_sales": float(current_sales),
            "projected_sales": float(projected_sales),
            "change": float(projected_sales - current_sales),
        })
    results.sort(key=lambda row: (row["price_list"].casefold(), row["mat_size"].casefold()))
    current_total = sum((Decimal(str(row["current_sales"])) for row in results), Decimal("0"))
    projected_total = sum((Decimal(str(row["projected_sales"])) for row in results), Decimal("0"))
    return {
        "analysis_type": "pricing",
        "analysis_request": str(request_text or "").strip(),
        "change_percent": change_percent,
        "value_column": selected_value,
        "quantity_column": selected_quantity,
        "price_list_column": selected_price_list,
        "mat_size_column": selected_mat_size,
        "value_mode": mode,
        "row_count": len(details),
        "excluded_rows": excluded_rows,
        "summary": {
            "current_sales": float(current_total),
            "projected_sales": float(projected_total),
            "change": float(projected_total - current_total),
            "change_percent": change_percent,
            "included_rows": len(details),
            "excluded_rows": len(excluded_rows),
        },
        "results": results,
        "details": details,
    }


def compare_company_rows(source_rows, companies, price_list_by_key=None, customer_metrics_by_key=None):
    price_list_by_key = price_list_by_key or {}
    customer_metrics_by_key = customer_metrics_by_key or {}
    prepared = []
    alias_index = {}
    for company in companies:
        aliases = _company_aliases(company)
        if not aliases:
            continue
        item = dict(company)
        item["aliases"] = aliases
        prepared.append(item)
        for alias in aliases:
            alias_index.setdefault(alias, []).append(item)
    all_aliases = list(alias_index)

    results = []
    for source in source_rows:
        normalized_source = normalize_company_name(source.get("source_company"))
        candidate_aliases = []
        if normalized_source in alias_index:
            candidate_aliases = [normalized_source]
        elif normalized_source:
            candidate_aliases = get_close_matches(normalized_source, all_aliases, n=8, cutoff=0.58)

        scored = []
        seen_keys = set()
        for alias in candidate_aliases:
            for company in alias_index.get(alias, []):
                company_key = str(company.get("primary_key") or company.get("filemaker_record_id") or company.get("company"))
                if company_key in seen_keys:
                    continue
                seen_keys.add(company_key)
                best_alias = max(company["aliases"], key=lambda value: _name_score(normalized_source, value))
                base_score = _name_score(normalized_source, best_alias)
                location_bonus, location_reasons = _location_bonus(source, company)
                score = min(100, base_score + location_bonus)
                scored.append((score, base_score, location_reasons, company))
        scored.sort(key=lambda item: (-item[0], str(item[3].get("company") or "").casefold()))

        best = scored[0] if scored else None
        ambiguous = bool(best and len(scored) > 1 and scored[1][0] >= best[0] - 2)
        if best and best[0] >= 90 and not ambiguous:
            match_status = "Matched"
        elif best and best[0] >= 76:
            match_status = "Possible match"
        else:
            match_status = "No match"
            best = None

        company = best[3] if best else {}
        primary_key = str(company.get("primary_key") or "").strip()
        price_list = str(company.get("price_list") or price_list_by_key.get(primary_key) or "").strip().upper()
        metrics = customer_metrics_by_key.get(primary_key) or {}
        reasons = []
        if best:
            reasons.append("exact normalized name" if best[1] == 100 else "similar company name")
            reasons.extend(best[2])
            if ambiguous:
                reasons.append("multiple close FileMaker matches")
        results.append({
            **source,
            "match_status": match_status,
            "matched_company": str(company.get("company") or ""),
            "customer_primary_key": primary_key,
            "matched_city": str(company.get("city") or ""),
            "matched_state": str(company.get("state") or ""),
            "matched_zip": str(company.get("zip_code") or ""),
            "price_list": price_list,
            "is_price_list_d": bool(best and match_status == "Matched" and price_list == "D"),
            "order_count": int(metrics.get("order_count") or 0) if best else 0,
            "total_spend": round(float(metrics.get("total_spend") or 0), 2) if best else 0,
            "confidence": best[0] if best else 0,
            "reason": ", ".join(reasons) if reasons else "No credible company match",
        })

    return results


def summarize_results(results):
    return {
        "uploaded": len(results),
        "matched": sum(row["match_status"] == "Matched" for row in results),
        "possible": sum(row["match_status"] == "Possible match" for row in results),
        "unmatched": sum(row["match_status"] == "No match" for row in results),
        "price_list_d": sum(bool(row.get("is_price_list_d")) for row in results),
    }


def save_analysis_run(payload, run_dir=DEFAULT_RUN_DIR):
    run_dir = Path(run_dir)
    run_dir.mkdir(parents=True, exist_ok=True)
    cleanup_analysis_runs(run_dir)
    token = secrets.token_urlsafe(24)
    path = run_dir / f"{token}.json"
    saved_payload = dict(payload)
    saved_payload["created_at"] = datetime.now().isoformat(timespec="seconds")
    path.write_text(json.dumps(saved_payload, ensure_ascii=False), encoding="utf-8")
    return token


def load_analysis_run(token, run_dir=DEFAULT_RUN_DIR):
    if not re.fullmatch(r"[A-Za-z0-9_-]{20,80}", str(token or "")):
        return None
    path = Path(run_dir) / f"{token}.json"
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None


def cleanup_analysis_runs(run_dir=DEFAULT_RUN_DIR):
    cutoff = datetime.now() - timedelta(hours=RUN_RETENTION_HOURS)
    for path in Path(run_dir).glob("*.json"):
        try:
            if datetime.fromtimestamp(path.stat().st_mtime) < cutoff:
                path.unlink()
        except OSError:
            continue


def build_results_workbook(payload):
    if payload.get("analysis_type") == "pricing":
        return build_pricing_workbook(payload)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Comparison"
    summary = payload.get("summary") or summarize_results(payload.get("results") or [])
    worksheet.append(["Company comparison"])
    worksheet.append(["Source file", payload.get("filename", "")])
    worksheet.append(["Company column", payload.get("company_column", "")])
    worksheet.append(["Analysis request", payload.get("analysis_request", "")])
    worksheet.append(["Order-history coverage", payload.get("coverage_note", "")])
    worksheet.append([])
    worksheet.append(["Uploaded", "Matched", "Possible matches", "No match", "Confirmed Price List D"])
    worksheet.append([summary["uploaded"], summary["matched"], summary["possible"], summary["unmatched"], summary["price_list_d"]])
    headers = [
        "Source row", "Uploaded company", "Uploaded city", "Uploaded state", "Uploaded postcode",
        "Match status", "Matched company", "Customer key", "Matched city", "Matched state",
        "Matched postcode", "Price list", "Price List D", "Order count", "Total spend", "Confidence", "Reason",
    ]
    worksheet.append(headers)
    for row in payload.get("results") or []:
        values = [
            row.get("source_row"), row.get("source_company"), row.get("source_city"), row.get("source_state"),
            row.get("source_zip"), row.get("match_status"), row.get("matched_company"),
            row.get("customer_primary_key"), row.get("matched_city"), row.get("matched_state"),
            row.get("matched_zip"), row.get("price_list"), "Yes" if row.get("is_price_list_d") else "No",
            row.get("order_count"), row.get("total_spend"), row.get("confidence"), row.get("reason"),
        ]
        worksheet.append([safe_excel_value(value) for value in values])

    navy = "17365D"
    blue = "D9EAF7"
    worksheet["A1"].font = Font(size=18, bold=True, color=navy)
    for cell in worksheet[7]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    for cell in worksheet[9]:
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(bold=True, color=navy)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    if worksheet.max_row >= 10:
        table = Table(displayName="CompanyComparison", ref=f"A9:Q{worksheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        worksheet.add_table(table)
    widths = [12, 30, 18, 16, 18, 18, 30, 18, 18, 16, 18, 12, 13, 13, 16, 12, 34]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[chr(64 + index)].width = width
    worksheet.freeze_panes = "A10"
    worksheet.auto_filter.ref = f"A9:Q{worksheet.max_row}"
    for cell in worksheet["O"][9:]:
        cell.number_format = '"$"#,##0.00'
    worksheet.sheet_view.showGridLines = False

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_pricing_workbook(payload):
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Pricing scenario"
    detail_sheet = workbook.create_sheet("Detail")
    summary = payload.get("summary") or {}
    results = payload.get("results") or []

    summary_sheet.append(["Pricing scenario"])
    summary_sheet.append(["Source file", payload.get("filename", "")])
    summary_sheet.append(["Analysis request", payload.get("analysis_request", "")])
    summary_sheet.append(["Value column", payload.get("value_column", "")])
    summary_sheet.append(["Quantity column", payload.get("quantity_column", "") or "Not used"])
    summary_sheet.append(["Value interpretation", "Unit value × quantity" if payload.get("value_mode") == "unit" else "Row sales value"])
    summary_sheet.append([])
    summary_sheet.append(["Current sales", "Projected sales", "Change", "Price change", "Included rows", "Excluded rows"])
    summary_sheet.append([
        summary.get("current_sales", 0), summary.get("projected_sales", 0), summary.get("change", 0),
        (summary.get("change_percent", 0) or 0) / 100, summary.get("included_rows", 0), summary.get("excluded_rows", 0),
    ])
    summary_sheet.append([])
    summary_sheet.append(["Price list", "Mat size", "Rows", "Quantity", "Current sales", "Projected sales", "Change"])
    for row in results:
        summary_sheet.append([
            safe_excel_value(row.get("price_list")), safe_excel_value(row.get("mat_size")), row.get("row_count"),
            row.get("quantity"), row.get("current_sales"), row.get("projected_sales"), row.get("change"),
        ])

    detail_sheet.append(["Source row", "Price list", "Mat size", "Quantity", "Input value", "Current sales", "Projected sales", "Change"])
    for row in payload.get("details") or []:
        detail_sheet.append([
            row.get("source_row"), safe_excel_value(row.get("price_list")), safe_excel_value(row.get("mat_size")),
            row.get("quantity"), row.get("input_value"), row.get("current_sales"), row.get("projected_sales"), row.get("change"),
        ])

    navy = "17365D"
    blue = "D9EAF7"
    summary_sheet["A1"].font = Font(size=18, bold=True, color=navy)
    for header_row in (8, 11):
        for cell in summary_sheet[header_row]:
            cell.fill = PatternFill("solid", fgColor=navy if header_row == 8 else blue)
            cell.font = Font(bold=True, color="FFFFFF" if header_row == 8 else navy)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in detail_sheet[1]:
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(bold=True, color=navy)
    if len(results) > 0:
        table = Table(displayName="PricingScenario", ref=f"A11:G{summary_sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        summary_sheet.add_table(table)
    if detail_sheet.max_row > 1:
        detail_table = Table(displayName="PricingDetail", ref=f"A1:H{detail_sheet.max_row}")
        detail_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        detail_sheet.add_table(detail_table)
    for sheet, widths in (
        (summary_sheet, [20, 20, 15, 15, 18, 18, 18]),
        (detail_sheet, [12, 18, 18, 14, 16, 18, 18, 18]),
    ):
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width
        sheet.sheet_view.showGridLines = False
    summary_sheet.freeze_panes = "A12"
    detail_sheet.freeze_panes = "A2"
    for row in summary_sheet.iter_rows(min_row=9, min_col=1, max_col=3):
        for cell in row:
            cell.number_format = '"$"#,##0.00'
    summary_sheet["D9"].number_format = "0.0%"
    for row in summary_sheet.iter_rows(min_row=12, min_col=5, max_col=7):
        for cell in row:
            cell.number_format = '"$"#,##0.00'
    for row in detail_sheet.iter_rows(min_row=2, min_col=5, max_col=8):
        for cell in row:
            cell.number_format = '"$"#,##0.00'

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def safe_excel_value(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value
